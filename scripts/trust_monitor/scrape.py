#!/usr/bin/env python3
"""14개 신탁사 공시 스크래핑, 전일 대비 신규 항목 탐지, xlsx 갱신."""

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook

HERE = Path(__file__).parent
REPORT_SHEET = "공시현황"
STATE_SHEET = "_state"
UA = {"User-Agent": "Mozilla/5.0 (compatible; trust-monitor/1.0)"}
DATE_RE = re.compile(r"(20\d{2})[.\-/년\s]+(\d{1,2})[.\-/월\s]+(\d{1,2})")
TITLE_KEYS = ("title", "subject", "name", "postTitle")
DATE_KEYS = ("date", "regDate", "createdAt", "postDate", "registDate", "createDate")


def norm_date(text):
    m = DATE_RE.search(text or "")
    return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if m else ""


def from_json(payload):
    """JSON 응답에서 제목/날짜를 가진 첫 레코드를 찾는다."""
    stack = [payload]
    while stack:
        node = stack.pop(0)
        if isinstance(node, dict):
            title = next((node[k] for k in TITLE_KEYS if isinstance(node.get(k), str)), None)
            if title:
                raw = next((str(node[k]) for k in DATE_KEYS if node.get(k)), "")
                return title.strip(), norm_date(raw)
            stack.extend(node.values())
        elif isinstance(node, list):
            stack.extend(node)
    return None


def from_html(html):
    """게시판 목록에서 날짜를 포함한 첫 행의 제목/날짜를 뽑는다."""
    # ponytail: 사이트별 셀렉터 대신 날짜 패턴 휴리스틱. 오탐 잦아지면 사별 셀렉터로 전환.
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "nav", "header", "footer"]):
        tag.decompose()
    for el in soup.find_all(["tr", "li"]):
        text = " ".join(el.get_text(" ", strip=True).split())
        if len(text) < 8 or not DATE_RE.search(text):
            continue
        anchor = el.find("a")
        title = " ".join(anchor.get_text(" ", strip=True).split()) if anchor else ""
        if not title:
            title = DATE_RE.sub("", text).strip(" |·-")
        if title:
            return title[:200], norm_date(text)
    return None


def fetch(url):
    if url.upper().startswith("POST "):
        raise ValueError("POST API 페이로드 미확인 — 실제 요청 형식 확인 필요")
    resp = requests.get(url, headers=UA, timeout=30)
    resp.raise_for_status()
    if "json" in resp.headers.get("content-type", ""):
        found = from_json(resp.json())
    else:
        found = from_html(resp.text)
    if not found:
        raise ValueError("목록에서 제목/날짜를 찾지 못함")
    return found


def fetch_kofia(code, cfg):
    """KOFIA DIS 경영공시에서 회사코드(cd) 기준 최신 항목."""
    template = cfg.get("payload")
    if not template:
        raise ValueError("KOFIA payload 미설정 — companies.json의 kofia.payload 필요")
    body = template.replace("{code}", code).replace("{sector}", cfg.get("sectorCode", ""))
    resp = requests.post(cfg["api"], data=body.encode("utf-8"), timeout=30,
                         headers={**UA, "Content-Type": "application/xml; charset=utf-8"})
    resp.raise_for_status()
    found = from_html(resp.text)
    if not found:
        raise ValueError("KOFIA 응답에서 제목/날짜를 찾지 못함")
    return found


def scrape(companies, kofia_cfg, kofia_fetcher=fetch_kofia):
    results = {}
    for c in companies:
        url = c.get("mgmt") or ""
        try:
            if not url:
                raise ValueError(c.get("status", "수집 URL 없음"))
            title, posted = fetch(url)
            results[c["co"]] = {"title": title, "posted": posted}
            continue
        except Exception as exc:
            primary_err = f"{type(exc).__name__}: {exc}"
        try:
            title, posted = kofia_fetcher(c["cd"], kofia_cfg)
            results[c["co"]] = {"title": title, "posted": posted, "via": "KOFIA"}
        except Exception as exc:
            results[c["co"]] = {"error": f"{primary_err} / KOFIA: {type(exc).__name__}: {exc}"}
    return results


def load_state(wb):
    if STATE_SHEET not in wb.sheetnames:
        return {}
    rows = wb[STATE_SHEET].iter_rows(min_row=2, values_only=True)
    return {r[0]: {"title": r[1] or "", "posted": r[2] or ""} for r in rows if r and r[0]}


def save_state(wb, results):
    if STATE_SHEET in wb.sheetnames:
        del wb[STATE_SHEET]
    ws = wb.create_sheet(STATE_SHEET)
    ws.sheet_state = "hidden"
    ws.append(["회사명", "최근제목", "최근일자", "갱신시각"])
    stamp = datetime.now().isoformat(timespec="seconds")
    for co, res in results.items():
        if "title" in res:
            ws.append([co, res["title"], res["posted"], stamp])


def write_report(wb, companies, notes, today):
    ws = wb[REPORT_SHEET] if REPORT_SHEET in wb.sheetnames else wb.create_sheet(REPORT_SHEET, 0)
    if ws["A1"].value is None:
        ws["A1"] = "회사명"
        for i, c in enumerate(companies, start=2):
            ws.cell(row=i, column=1, value=c["co"])
    header = [ws.cell(row=1, column=col).value for col in range(2, ws.max_column + 1)]
    col = header.index(today) + 2 if today in header else ws.max_column + 1
    ws.cell(row=1, column=col, value=today)
    for row in range(2, ws.max_row + 1):
        co = ws.cell(row=row, column=1).value
        ws.cell(row=row, column=col, value=notes.get(co, ""))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", help="기존 xlsx 경로 (없으면 새로 생성)")
    ap.add_argument("--out", required=True, help="갱신된 xlsx 저장 경로")
    ap.add_argument("--companies", default=str(HERE / "companies.json"))
    args = ap.parse_args()

    source = json.loads(Path(args.companies).read_text(encoding="utf-8"))
    companies, kofia_cfg = source["companies"], source.get("kofia", {})
    today = date.today().isoformat()

    if args.workbook and Path(args.workbook).exists():
        wb = load_workbook(args.workbook)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    previous = load_state(wb)

    results = scrape(companies, kofia_cfg)
    notes, changed, errors = {}, [], []
    for co, res in results.items():
        if "error" in res:
            notes[co] = f"수집실패: {res['error'][:80]}"
            errors.append(co)
        elif previous.get(co, {}).get("title") != res["title"]:
            first_run = co not in previous
            source_tag = " (KOFIA)" if res.get("via") else ""
            notes[co] = ("최초수집: " if first_run else "신규공시: ") + res["title"][:120] + source_tag
            if not first_run:
                changed.append(co)

    write_report(wb, companies, notes, today)
    save_state(wb, results)
    wb.save(args.out)

    json.dump(
        {"date": today, "changed": changed, "errors": errors,
         "notes": {k: v for k, v in notes.items() if v}, "out": args.out},
        sys.stdout, ensure_ascii=False, indent=2,
    )


if __name__ == "__main__":
    main()
