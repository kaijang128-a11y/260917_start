"""scrape.py 워크북/파싱 로직 자체검증. 네트워크 불필요."""

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from scrape import from_html, from_json, load_state, norm_date, save_state, write_report

COMPANIES = [{"co": "우리자산신탁"}, {"co": "하나자산신탁"}]

LIST_HTML = """
<table><tr><th>제목</th><th>등록일</th></tr>
<tr><td><a href="/view/1">2026년 반기 경영공시</a></td><td>2026.09.15</td></tr>
<tr><td><a href="/view/2">2026년 1분기 경영공시</a></td><td>2026.05.14</td></tr></table>
"""


def test_norm_date():
    assert norm_date("2026.09.15") == "2026-09-15"
    assert norm_date("2026년 9월 3일") == "2026-09-03"
    assert norm_date("제목만 있음") == ""


def test_from_html():
    title, posted = from_html(LIST_HTML)
    assert title == "2026년 반기 경영공시", title
    assert posted == "2026-09-15", posted


def test_from_json():
    payload = {"data": {"content": [{"title": "수시공시 제출", "createdAt": "2026-09-16T10:00:00"}]}}
    assert from_json(payload) == ("수시공시 제출", "2026-09-16")


def test_report_appends_column_per_date(tmp="/tmp/tm_selfcheck.xlsx"):
    wb = Workbook()
    wb.remove(wb.active)
    write_report(wb, COMPANIES, {"우리자산신탁": "최초수집: A"}, "2026-09-17")
    save_state(wb, {"우리자산신탁": {"title": "A", "posted": "2026-09-15"}})
    wb.save(tmp)

    wb2 = load_workbook(tmp)
    assert load_state(wb2)["우리자산신탁"]["title"] == "A"
    write_report(wb2, COMPANIES, {"하나자산신탁": "신규공시: B"}, "2026-09-18")
    ws = wb2["공시현황"]
    assert [ws.cell(row=1, column=c).value for c in (1, 2, 3)] == ["회사명", "2026-09-17", "2026-09-18"]
    assert ws.cell(row=2, column=2).value == "최초수집: A"
    assert ws.cell(row=3, column=3).value == "신규공시: B"
    assert wb2["_state"].sheet_state == "hidden"


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_"):
            fn()
            print(f"ok {name}")
